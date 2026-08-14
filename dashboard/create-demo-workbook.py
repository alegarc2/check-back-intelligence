#!/usr/bin/env python3
"""Create a sanitized demo Check Back workbook for public GitHub Pages hosting."""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from src.workbook.check_back_template import CHECK_BACK_HEADERS  # noqa: E402

HDR_FILL = PatternFill(start_color="0D1526", end_color="0D1526", fill_type="solid")
HDR_FONT = Font(bold=True, color="94A3B8", size=10)

DEMO_ROWS = [
    {
        "Opportunity Name": "Demo School District (sample)",
        "SL2": "US PS Market",
        "TCV $": 850000,
        "Partner": "Demo Partner LLC",
        "Sub #": "Sub900001",
        "Provisioned/Entitled Lic Calling": "Professional 283/4,200; Workspace 1,200/1,750",
        "Active Lic Calling": "420",
        "(G/Y/R)": "G",
        "Customer org id": "707f1259-feb5-4ea2-9a06-05b36402f6cf",
        "Platforms": "Webex",
    },
    {
        "Opportunity Name": "Demo Health System (sample)",
        "SL2": "US Commercial",
        "TCV $": 1200000,
        "Partner": "Demo Partner LLC",
        "Sub #": "Sub900002",
        "Provisioned/Entitled Lic Calling": "Professional 1,326/1,860; Workspace 374/775",
        "Active Lic Calling": "890",
        "(G/Y/R)": "Y",
        "Customer org id": "a1b2c3d4-e5f6-7890-abcd-ef1234567890",
        "Platforms": "Webex",
    },
    {
        "Opportunity Name": "Demo Manufacturing Co (sample)",
        "SL2": "EMEA__UKI",
        "TCV $": 640000,
        "Partner": "Demo Partner GmbH",
        "Sub #": "Sub900003",
        "Provisioned/Entitled Lic Calling": "Professional 129/1,800; Workspace 57/750",
        "Active Lic Calling": "95",
        "(G/Y/R)": "R",
        "Customer org id": "b2c3d4e5-f6a7-8901-bcde-f23456789012",
        "Platforms": "Webex",
    },
]


def main() -> None:
    dest = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "samples" / "check_back_demo.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customer Data"
    for col, name in enumerate(CHECK_BACK_HEADERS, start=1):
        cell = ws.cell(1, col, value=name)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r_idx, row in enumerate(DEMO_ROWS, start=2):
        for col, name in enumerate(CHECK_BACK_HEADERS, start=1):
            val = row.get(name)
            if val not in (None, ""):
                ws.cell(r_idx, col, value=val)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    print(f"Wrote {dest}")


if __name__ == "__main__":
    main()
